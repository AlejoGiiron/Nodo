# ============================================================================
# VERIFICACION DEL HISTORICO v3 — INDEPENDIENTE DEL CARGADOR.
#
# 🔴 POR QUE ES PYTHON Y EL CARGADOR ES NODE. El cargador lee el Excel con
#    `exceljs` y escribe con `supabase-js`. Este lo lee con `openpyxl` y consulta
#    por HTTP crudo contra PostgREST. **Comparar la base contra la tabla que el
#    propio cargador leyo verificaria la base contra su propio typo.**
#
# 🔴 Y EL CRITERIO INCLUYE QUE SE VEA POR LOS CAMINOS DEL PRODUCTO, no solo que
#    los numeros cierren — la leccion de la numeracion: nueve verificaciones en
#    verde y las 30 ventas INVISIBLES en el Historial por falta de `order_number`.
#
# USO: MPE=<correo> MPP=<clave> python scripts/verificar-historico-v3.py --sede-id <uuid>
# ============================================================================
import io, json, os, re, sys, unicodedata, urllib.request
from collections import defaultdict
import openpyxl
from decimal import Decimal, ROUND_HALF_UP

ARCHIVO = 'docs/Control Mp 3.xlsx'
FANTASMA = '4b984fb3-8278-4370-94f4-ba98e8693298'   # residuo previo, no es del historico

def arg(n):
    a = sys.argv
    return a[a.index('--' + n) + 1] if '--' + n in a else None

SEDE = arg('sede-id')
if not SEDE: sys.exit('falta --sede-id')
MPE, MPP = os.environ.get('MPE'), os.environ.get('MPP')
if not (MPE and MPP): sys.exit('faltan MPE/MPP')

env = io.open('.env', encoding='utf-8').read()
de = lambda k: (re.search('^' + k + '=(.*)$', env, re.M) or [None, ''])[1].strip().strip('"\'')
URL, ANON = de('VITE_NODO_SUPABASE_URL'), de('VITE_NODO_SUPABASE_ANON_KEY')

def http(ruta, token=None, datos=None, prefer=None):
    req = urllib.request.Request(URL + ruta, method='POST' if datos else 'GET')
    req.add_header('apikey', ANON)
    req.add_header('Authorization', 'Bearer ' + (token or ANON))
    if datos: req.add_header('Content-Type', 'application/json')
    if prefer: req.add_header('Prefer', prefer)
    cuerpo = json.dumps(datos).encode() if datos else None
    with urllib.request.urlopen(req, cuerpo) as r:
        return json.loads(r.read().decode()), dict(r.headers)

tok = http('/auth/v1/token?grant_type=password', datos={'email': MPE, 'password': MPP})[0]['access_token']

def rest(tabla, query):
    """Lee TODAS las filas paginando. 🔴 PostgREST corta en 1000 sin avisar."""
    filas, desde = [], 0
    while True:
        req = urllib.request.Request(URL + '/rest/v1/' + tabla + '?' + query)
        req.add_header('apikey', ANON); req.add_header('Authorization', 'Bearer ' + tok)
        req.add_header('Range', '%d-%d' % (desde, desde + 999))
        req.add_header('Prefer', 'count=exact')
        with urllib.request.urlopen(req) as r:
            d = json.loads(r.read().decode()); tot = int(r.headers['Content-Range'].split('/')[-1])
        filas += d; desde += len(d)
        if desde >= tot or not d: return filas, tot

def plano(x):
    s = unicodedata.normalize('NFD', str(x or '').strip().lower())
    return re.sub(r'\s+', ' ', ''.join(c for c in s if unicodedata.category(c) != 'Mn'))

FALLAS, N = [], 0
def chk(nombre, esperado, obtenido, nota=''):
    global N
    N += 1
    ok = esperado == obtenido
    print('  %s %-52s esperado %-16s obtenido %s%s'
          % ('✅' if ok else '🔴', nombre, str(esperado), str(obtenido), ('  · ' + nota) if nota else ''))
    if not ok: FALLAS.append(nombre)

# ════════════════════════════════════════════════════════════════════════════
# § 1 · EL ARCHIVO, leido con openpyxl (NO con el parser del cargador)
# ════════════════════════════════════════════════════════════════════════════
wb = openpyxl.load_workbook(ARCHIVO, data_only=True)
def cols(h):
    ws = wb[h]; return ws, {str(ws.cell(1, i).value or '').strip().lower(): i
                            for i in range(1, ws.max_column + 1)}
def col(c, sub):
    k = [v for k, v in c.items() if k == sub or sub in k]
    if not k: sys.exit('no encuentro la columna «%s»' % sub)
    return k[0]

wc, cc = cols('Compra de inventario')
compras = []
for r in range(2, wc.max_row + 1):
    f, p = wc.cell(r, cc['fecha']).value, wc.cell(r, cc['producto']).value
    q, cu = wc.cell(r, col(cc, 'unidades')).value, wc.cell(r, col(cc, 'costo unitario')).value
    pv = wc.cell(r, col(cc, 'proveedor')).value
    if f and p and isinstance(q, (int, float)) and isinstance(cu, (int, float)):
        compras.append({'f': str(f)[:10], 'p': plano(p), 'q': q, 'cu': cu, 'prov': plano(pv)})

wv, cv = cols('Ventas Diarias')
ventas = []
for r in range(2, wv.max_row + 1):
    f, p = wv.cell(r, cv['fecha']).value, wv.cell(r, cv['producto']).value
    q = wv.cell(r, col(cv, 'cantidad vendida')).value
    pr = wv.cell(r, col(cv, 'precio real')).value
    if f and p and isinstance(q, (int, float)):
        ventas.append({'f': str(f)[:10], 'p': plano(p), 'q': q,
                       'pr': pr if isinstance(pr, (int, float)) else 0,
                       'tipo': str(wv.cell(r, col(cv, 'tipo de pago')).value or '').upper(),
                       'met': str(wv.cell(r, col(cv, 'metodo de pago')).value or '').upper()})

wg, cg = cols('Gasto')
gastos = [ {'f': str(wg.cell(r, cg['fecha']).value)[:10], 'v': wg.cell(r, col(cg, 'valor')).value}
           for r in range(2, wg.max_row + 1)
           if wg.cell(r, cg['fecha']).value and isinstance(wg.cell(r, col(cg, 'valor')).value, (int, float)) ]

VENDIDO   = sum(v['q'] * v['pr'] for v in ventas)
COMPRADO  = sum(c['q'] * c['cu'] for c in compras)
GASTOS    = sum(g['v'] for g in gastos)
CREDITO   = sum(v['q'] * v['pr'] for v in ventas if v['tipo'].startswith('CRED'))
ABONOS    = sum(int(re.search(r'ABONO\s+([\d.]+)\s*MIL', v['met']).group(1).replace('.', '')) * 1000
                for v in ventas if re.match(r'^ABONO\s', v['met']))
DIAS      = sorted({x['f'] for x in compras} | {x['f'] for x in ventas} | {g['f'] for g in gastos})
COBRADAS  = [v for v in ventas if not v['tipo'].startswith('CRED')]

print('══ EL ARCHIVO (parser: openpyxl) ══')
print('   %d dias · %d compras · %d ventas · %d gastos' % (len(DIAS), len(compras), len(ventas), len(gastos)))
print('   comprado %s · vendido %s · gastos %s · credito %s · abonos %s'
      % tuple(format(round(x), ',d') for x in (COMPRADO, VENDIDO, GASTOS, CREDITO, ABONOS)))

# ════════════════════════════════════════════════════════════════════════════
# § 2 · LA BASE, por HTTP crudo
# ════════════════════════════════════════════════════════════════════════════
ords, _   = rest('orders', 'sede_id=eq.%s&select=id,order_number,total,created_at,payment_status,customer_id,canal,cancelled_at' % SEDE)
ords      = [o for o in ords if o['id'] != FANTASMA]
oid       = {o['id'] for o in ords}
items, _  = rest('order_items', 'select=id,order_id,product_id,qty,unit_price,unit_cost,nivel_aplicado')
items     = [i for i in items if i['order_id'] in oid]
pays, _   = rest('payments', 'sede_id=eq.%s&select=order_id,method,amount' % SEDE)
dpays, _  = rest('debt_payments', 'select=order_id,amount,payment_method')
dpays     = [d for d in dpays if d['order_id'] in oid]
facs, _   = rest('purchase_invoices', 'sede_id=eq.%s&select=id,total,document_date,supplier_id' % SEDE)
movs, _   = rest('cash_movements', 'sede_id=eq.%s&select=amount,categoria,subcategoria,document_date' % SEDE)
jors, _   = rest('jornadas', 'sede_id=eq.%s&select=id,opened_at,closed_at,closing_amount,expected_amount,difference' % SEDE)
prods, _  = rest('products', 'sede_id=eq.%s&select=id,name,stock_qty,cost_price,is_active,price' % SEDE)
sups, _   = rest('suppliers', 'sede_id=eq.%s&select=id,name' % SEDE)

print('\n══ A · LOS CONJUNTOS Y LAS SUMAS ══')
chk('ordenes del historico', len(ventas), len(ords))
chk('lineas de venta', len(ventas), len(items))
chk('vendido (suma de qty x unit_price)', round(VENDIDO), round(sum(i['qty'] * float(i['unit_price']) for i in items)))
chk('total derivado por el servidor', round(VENDIDO), round(sum(float(o['total']) for o in ords)))
chk('pagos (las cobradas)', len(COBRADAS), len(pays))
chk('cobrado', round(VENDIDO - CREDITO), round(sum(float(p['amount']) for p in pays)))
chk('ordenes a credito', len(ventas) - len(COBRADAS), len([o for o in ords if o['payment_status'] != 'paid']))
chk('abonos', 2, len(dpays))
chk('abonado', ABONOS, round(sum(float(d['amount']) for d in dpays)))
chk('SALDO DE CARTERA', round(CREDITO - ABONOS), round(CREDITO - sum(float(d['amount']) for d in dpays)))
chk('facturas de compra', len(compras), len(facs))
# 🔴 EL COMPRADO SE ASEVERA CON DOS SUMAS, NO CON UNA — y la diferencia entre
#    ellas es un HECHO DEL ESQUEMA, no un error de carga:
#    `purchase_invoice_items.unit_cost` es `numeric(12,2)`, asi que el costo
#    unitario se REDONDEA AL GUARDARSE y el subtotal sale del ya redondeado.
#    Catorce lineas del archivo traen costos de 3+ decimales (son formulas de
#    Excel), y diez de ellas aportan 0,09 sobre 13,1 millones.
#    ⚠️ La version anterior aseveraba solo la suma cruda y daba ROJO sobre datos
#    CORRECTOS — el modo de fallo que este proyecto ya midio: la reaccion natural
#    ante una suma que no cierra es tocar los datos hasta que cierre.
COMPRADO_GUARDADO = sum(
    float((Decimal(str(c['q'])) * Decimal(str(c['cu'])).quantize(Decimal('0.01'), ROUND_HALF_UP))
          .quantize(Decimal('0.01'), ROUND_HALF_UP)) for c in compras)
enBase = round(sum(float(f['total']) for f in facs), 2)
chk('comprado (tal como el esquema lo GUARDA)', round(COMPRADO_GUARDADO, 2), enBase)
chk('la diferencia con la suma cruda del archivo es solo redondeo de columna',
    round(COMPRADO_GUARDADO - COMPRADO, 2), round(enBase - COMPRADO, 2),
    'numeric(12,2) sobre 14 costos con 3+ decimales')
gas = [m for m in movs if m['categoria'] == 'gasto']
chk('gastos', len(gastos), len(gas))
chk('gastado', round(GASTOS), round(sum(float(m['amount']) for m in gas)))
chk('proveedores (VENOM/Venom colapsados)', len({c['prov'] for c in compras}), len(sups))
chk('canal mostrador en todas', len(ventas), len([o for o in ords if o['canal'] == 'mostrador']))
chk('ninguna anulada', 0, len([o for o in ords if o['cancelled_at']]))

metodo = defaultdict(int)
for p in pays: metodo[p['method']] += 1
arch = defaultdict(int)
for v in COBRADAS: arch[v['met']] += 1
chk('metodo TRANSFERENCIA -> transfer', arch['TRANSFERENCIA'], metodo['transfer'])
chk('metodo EFECTIVO -> cash', arch['EFECTIVO'], metodo['cash'])

print('\n══ B · LA MEDICION QUE DECIDIO EL ORDEN DE CARGA ══')
nulos = [i for i in items if i['unit_cost'] is None]
chk('lineas con unit_cost NULO (medido: 0 de 110)', 0, len(nulos),
    'el orden contrario daba 39 y 49,8% del vendido')
chk('lineas con unit_cost en CERO', 0, len([i for i in items if i['unit_cost'] is not None and float(i['unit_cost']) == 0]),
    'un cero seria un margen del 100% falso')
chk('nivel_aplicado NULO en todas (son anteriores a las listas)', len(ventas),
    len([i for i in items if i['nivel_aplicado'] is None]))

print('\n══ C · EL INVENTARIO SALE DEL HISTORICO ══')
esp = defaultdict(int)
for c in compras: esp[c['p']] += c['q']
for v in ventas:  esp[v['p']] -= v['q']
porNombre = {plano(p['name']): p for p in prods}
malos, sinFila = [], []
for n, q in esp.items():
    p = porNombre.get(n)
    if not p: sinFila.append(n); continue
    if (p['stock_qty'] or 0) != q: malos.append((n, q, p['stock_qty']))
chk('productos del archivo con fila en la base', 0, len(sinFila), str(sinFila[:4]))
chk('stock = Sum(compras) - Sum(ventas), producto por producto', 0, len(malos),
    str(malos[:3]) if malos else 'sobre %d productos movidos' % len(esp))
chk('productos en negativo', 0, len([p for p in prods if (p['stock_qty'] or 0) < 0]))
chk('productos con existencia', len([q for q in esp.values() if q > 0]),
    len([p for p in prods if (p['stock_qty'] or 0) > 0]))
chk('unidades en existencia', sum(q for q in esp.values() if q > 0),
    sum(p['stock_qty'] for p in prods if (p['stock_qty'] or 0) > 0))

print('\n══ D · LAS JORNADAS ══')
# 🔴 DOS BLOQUES, no un filtro: las del HISTORICO se comparan contra los 15 dias
#    del archivo, y las ADMINISTRATIVAS —dias en que tocamos la sede sin que ella
#    operara— se cuentan y se nombran aparte. La version anterior tenia un solo
#    numero y se ponia roja sin decir cual de las dos cosas habia cambiado.
def diaBogotaPre(ts):
    from datetime import datetime, timedelta, timezone
    return datetime.fromisoformat(ts.replace('Z', '+00:00')).astimezone(timezone(timedelta(hours=-5))).date().isoformat()
delHist = [j for j in jors if diaBogotaPre(j['opened_at']) in DIAS]
admin   = [j for j in jors if diaBogotaPre(j['opened_at']) not in DIAS]
if admin:
    print('  ⚠️ %d jornada(s) ADMINISTRATIVA(S), fuera de los 15 dias del archivo: %s'
          % (len(admin), ', '.join(sorted(diaBogotaPre(j['opened_at']) for j in admin))))
    print('     (dias en que tocamos la sede; no son operacion suya y no se suman al historico)')
chk('jornadas del historico', len(DIAS), len(delHist))
chk('ninguna quedo ABIERTA (bloquearia la sede entera)', 0, len([j for j in jors if not j['closed_at']]))
jors = delHist
chk('todas cerradas', len(DIAS), len([j for j in jors if j['closed_at']]))
def diaBogota(ts):
    # 🔴 R7: la frontera de dia se calcula en America/Bogota, no sobre el UTC crudo
    from datetime import datetime, timedelta, timezone
    return (datetime.fromisoformat(ts.replace('Z', '+00:00')).astimezone(timezone(timedelta(hours=-5)))).date().isoformat()
mismodia = [j for j in jors if diaBogota(j['opened_at']) == diaBogota(j['closed_at'])]
chk('cierre en SU MISMO dia de Bogota', len(DIAS), len(mismodia))
chk('los 15 dias del archivo, uno por uno', DIAS, sorted(diaBogota(j['opened_at']) for j in jors))
chk('SIN arqueo (closing_amount nulo en las 15)', len(DIAS), len([j for j in jors if j['closing_amount'] is None]),
    'un cero seria un arqueo falso persistido')

print('\n══ E · QUE SE VEA POR LOS CAMINOS DEL PRODUCTO ══')
# HISTORIAL: ordena por order_number. Sin numero, la venta no aparece.
sinNum = [o for o in ords if o['order_number'] is None]
nums = [o['order_number'] for o in ords if o['order_number'] is not None]
chk('HISTORIAL · ventas sin order_number (serian invisibles)', 0, len(sinNum))
chk('HISTORIAL · numeros duplicados', 0, len(nums) - len(set(nums)))
# los bloques por dia van en orden (dentro de un dia el timestamp empata: R7)
porDia = defaultdict(list)
for o in ords:
    if o['order_number']: porDia[diaBogota(o['created_at'])].append(o['order_number'])
desorden = [d for i, d in enumerate(sorted(porDia)[:-1])
            if max(porDia[d]) > min(porDia[sorted(porDia)[i + 1]])]
chk('HISTORIAL · los bloques por dia van en orden', 0, len(desorden), str(desorden))
# CATALOGO
chk('CATALOGO · productos activos', 62, len([p for p in prods if p['is_active']]))
chk('CATALOGO · productos con costo (ya no nulo)', len({c['p'] for c in compras}),
    len([p for p in prods if p['cost_price'] is not None]))
# ── CARTERA ────────────────────────────────────────────────────────────────
# 🔴 SE MIDE CON LA CONSULTA DEL PRODUCTO, COPIADA Y SIN FILTROS PROPIOS.
#    `getDebts` (src/lib/supabase-helpers.ts) filtra exactamente esto. La version
#    anterior excluia la orden fantasma POR UUID y reportaba «9 ✅» sobre una
#    pantalla que mostraba 10: verificaba el conjunto que habia elegido, no el
#    que la clienta ve. Es la leccion de las 30 ventas invisibles INVERTIDA —
#    alla los datos estaban y no se veian; aca se ve algo que no se contaba.
cart, _ = rest('orders', ('sede_id=eq.%s&payment_status=in.(pending,partial)'
                          '&cancelled_at=is.null&select=id,order_number,total,customer_id,'
                          'debt_payments(amount)') % SEDE)
saldoCart = sum(float(o['total']) - sum(float(d['amount']) for d in (o.get('debt_payments') or []))
                for o in cart)
chk('CARTERA · filas que muestra la pantalla', len(ventas) - len(COBRADAS), len(cart),
    'numeros: ' + ', '.join(str(o['order_number']) for o in sorted(cart, key=lambda x: x['order_number'] or 0)))
chk('CARTERA · saldo total', round(CREDITO - ABONOS), round(saldoCart))
chk('CARTERA · clientes distintos con deuda', 7, len({o['customer_id'] for o in cart}),
    'los nombres no se imprimen: son PII')
# y el bloque APARTE, que contesta la OTRA pregunta: «¿se cargo bien?»
pend = [o for o in ords if o['payment_status'] != 'paid']
chk('(la carga, aparte) ordenes del historico con saldo', len(ventas) - len(COBRADAS), len(pend),
    'este SI excluye lo ajeno: es otra pregunta y va en otra linea')
# INVENTARIO
chk('INVENTARIO · referencias con existencia > 0', len([q for q in esp.values() if q > 0]),
    len([p for p in prods if (p['stock_qty'] or 0) > 0]))

print('\n' + '═' * 78)
print('  %d verificaciones · %d en verde · %d en rojo' % (N, N - len(FALLAS), len(FALLAS)))
if FALLAS:
    print('  🔴 ROJAS: ' + ' | '.join(FALLAS))
    sys.exit(1)
print('  ✅ TODO CIERRA contra el archivo, con otro parser, y se ve por los cuatro caminos.')
