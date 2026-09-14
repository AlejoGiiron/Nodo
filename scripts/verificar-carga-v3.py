# ============================================================================
# VERIFICACION DE LA CARGA v3 — INDEPENDIENTE DEL CARGADOR.
#
# 🔴 POR QUE ESTE ARCHIVO ES DE PYTHON Y EL CARGADOR DE NODE. El cargador lee el
#    Excel con `exceljs`; este lo lee con `openpyxl`. **Comparar la base contra
#    la tabla que el propio cargador leyo verificaria la base contra su propio
#    typo** — dos lecturas del mismo parser coinciden aunque las dos esten mal.
#    La lectura de la base va por HTTP contra PostgREST, sin el cliente de JS.
#
# 🔴 Y EL CRITERIO DE ACEPTACION INCLUYE QUE LOS DATOS SE VEAN POR LOS CAMINOS
#    DEL PRODUCTO, no solo que los numeros cierren. Es la leccion de la
#    numeracion del historico: nueve verificaciones en verde y las 30 ventas
#    INVISIBLES en el Historial porque les faltaba `order_number`. Aca las
#    consultas son las de `src/lib/supabase-helpers.ts`, copiadas, no
#    equivalentes.
#
# USO:  MPE=<correo> MPP=<clave> python scripts/verificar-carga-v3.py --sede-id <uuid>
# ============================================================================
import io, json, os, re, sys, unicodedata, urllib.request
from collections import Counter
import openpyxl

ARCHIVO = 'docs/Control Mp 3.xlsx'

def arg(n):
    a = sys.argv
    return a[a.index('--' + n) + 1] if '--' + n in a else None

SEDE = arg('sede-id')
if not SEDE: sys.exit('falta --sede-id')
MPE, MPP = os.environ.get('MPE'), os.environ.get('MPP')
if not (MPE and MPP): sys.exit('faltan MPE/MPP')

env = io.open('.env', encoding='utf-8').read()
de = lambda k: (re.search('^' + k + '=(.*)$', env, re.M) or [None, ''])[1].strip().strip('"\'')
URL, ANON = de('VITE_NODO_SUPABASE_URL'), de('VITE_NODO_SUPABASE_ANON_KEY')

def http(ruta, token=None, datos=None, prefer=None):
    req = urllib.request.Request(URL + ruta, method='POST' if datos else 'GET')
    req.add_header('apikey', ANON)
    req.add_header('Authorization', 'Bearer ' + (token or ANON))
    if datos: req.add_header('Content-Type', 'application/json')
    if prefer: req.add_header('Prefer', prefer)
    cuerpo = json.dumps(datos).encode() if datos else None
    with urllib.request.urlopen(req, cuerpo) as r:
        return json.loads(r.read().decode()), dict(r.headers)

tok = http('/auth/v1/token?grant_type=password', datos={'email': MPE, 'password': MPP})[0]['access_token']

def rest(tabla, query, contar=False):
    """Lee por PostgREST. `contar` pide el total exacto en la cabecera."""
    d, h = http('/rest/v1/' + tabla + '?' + query, tok, prefer='count=exact' if contar else None)
    total = None
    if contar and 'Content-Range' in h:
        total = int(h['Content-Range'].split('/')[-1])
    return d, total

def plano(x):
    s = unicodedata.normalize('NFD', str(x or '').strip().lower())
    return re.sub(r'\s+', ' ', ''.join(c for c in s if unicodedata.category(c) != 'Mn'))

# ── EL ARCHIVO, leido con openpyxl ──────────────────────────────────────────
wb = openpyxl.load_workbook(ARCHIVO, data_only=True)
def cols(h):
    ws = wb[h]
    return ws, {str(ws.cell(1, c).value or '').strip().lower(): c for c in range(1, ws.max_column + 1)}

wsP, cP = cols('Productos')
maestro = {}
for r in range(2, wsP.max_row + 1):
    n = wsP.cell(r, cP['producto']).value
    if n: maestro[plano(n)] = {'nombre': str(n).strip(),
                               'codigo': str(wsP.cell(r, cP['item']).value or '').strip(),
                               'cat': plano(wsP.cell(r, cP['categoria']).value)}

wsC, cC = cols('Compra de inventario')
niv = {}
for r in range(2, wsC.max_row + 1):
    n = wsC.cell(r, cC['producto']).value
    if not n: continue
    v = {k: wsC.cell(r, cC[k.lower()]).value for k in ['L0', 'L1', 'L2', 'L3', 'L4']}
    if not all(isinstance(x, (int, float)) and x > 0 for x in v.values()): continue
    f = wsC.cell(r, cC['fecha']).value
    k = plano(n)
    if k not in niv or (f and niv[k]['f'] and f >= niv[k]['f']): niv[k] = {'f': f, 'v': v}

wsV, cV = cols('Ventas Diarias')
clientes = {str(wsV.cell(r, cV['cliente']).value).strip()
            for r in range(2, wsV.max_row + 1) if wsV.cell(r, cV['cliente']).value}

print('══ EL ARCHIVO (parser: openpyxl) ══')
print('   productos: %d · con los 5 niveles: %d · clientes: %d · categorias: %d'
      % (len(maestro), len(niv), len(clientes), len({m['cat'] for m in maestro.values()})))

# ── SUMAS DE CONTROL del archivo ────────────────────────────────────────────
sumaL1 = sum(round(niv[k]['v']['L1']) for k in maestro)
suma310 = sum(round(niv[k]['v'][x]) for k in maestro for x in ['L0', 'L1', 'L2', 'L3', 'L4'])
porNivel = {x: sum(round(niv[k]['v'][x]) for k in maestro) for x in ['L0', 'L1', 'L2', 'L3', 'L4']}

fallos = []
def chequear(que, obtenido, esperado):
    ok = obtenido == esperado
    print(('   ✅ ' if ok else '   🔴 ') + que.ljust(46) + str(obtenido) + ('' if ok else '   esperaba ' + str(esperado)))
    if not ok: fallos.append(que)

# ── 1 · LOS NUMEROS, contra la base ─────────────────────────────────────────
print('\n══ 1 · CONTEOS Y SUMAS · base contra archivo ══')
_, nCat = rest('categories', 'sede_id=eq.%s&select=id' % SEDE, contar=True)
chequear('categorias', nCat, len({m['cat'] for m in maestro.values()}))
_, nPro = rest('products', 'sede_id=eq.%s&select=id' % SEDE, contar=True)
chequear('productos', nPro, len(maestro))
_, nPre = rest('product_prices', 'sede_id=eq.%s&select=product_id' % SEDE, contar=True)
chequear('filas de precio', nPre, len(maestro) * 5)
_, nCli = rest('customers', 'sede_id=eq.%s&select=id' % SEDE, contar=True)
chequear('clientes', nCli, len(clientes))

prods, _ = rest('products', 'sede_id=eq.%s&select=id,name,codigo,unidad,price,category_id&limit=1000' % SEDE)
chequear('SUMA de los %d precios L1' % len(maestro), sum(int(p['price']) for p in prods), sumaL1)

precios, _ = rest('product_prices', 'sede_id=eq.%s&select=product_id,nivel,precio&limit=2000' % SEDE)
chequear('SUMA de las %d filas de precio' % (len(maestro) * 5), sum(int(float(p['precio'])) for p in precios), suma310)
pn = Counter()
for p in precios: pn['L%d' % p['nivel']] += int(float(p['precio']))
for x in ['L0', 'L1', 'L2', 'L3', 'L4']:
    chequear('  suma de ' + x, pn[x], porNivel[x])

# ── 2 · FILA POR FILA ───────────────────────────────────────────────────────
print('\n══ 2 · FILA POR FILA · nombre, codigo, precio ══')
porNombre = {plano(p['name']): p for p in prods}
malCodigo = [k for k in maestro if porNombre.get(k, {}).get('codigo') != maestro[k]['codigo']]
malPrecio = [k for k in maestro if k in porNombre and int(porNombre[k]['price']) != round(niv[k]['v']['L1'])]
faltan = [k for k in maestro if k not in porNombre]
sobran = [k for k in porNombre if k not in maestro]
chequear('productos del archivo que NO estan', len(faltan), 0)
chequear('productos en la base que NO estan en el archivo', len(sobran), 0)
chequear('con codigo distinto al del archivo', len(malCodigo), 0)
chequear('con precio distinto a su L1', len(malPrecio), 0)
for k in (faltan + sobran + malCodigo + malPrecio)[:8]: print('        ' + k[:52])
chequear('sin unidad', sum(1 for p in prods if not p['unidad']), 0)
# los codigos repetidos de su archivo tienen que estar REPETIDOS, no inventados
repArchivo = Counter(m['codigo'] for m in maestro.values())
repBase = Counter(p['codigo'] for p in prods)
chequear('codigos repetidos igual que en el archivo',
         sorted((c, n) for c, n in repBase.items() if n > 1),
         sorted((c, n) for c, n in repArchivo.items() if n > 1))

# ── 3 · QUE SE VEAN POR LOS CAMINOS DEL PRODUCTO ────────────────────────────
print('\n══ 3 · VISIBILIDAD · las consultas del producto, no equivalentes ══')
# getProducts(sedeId) de src/lib/supabase-helpers.ts
cat, _ = rest('products',
  'select=*,categories(id,name,color),product_prices(nivel,precio)&sede_id=eq.%s&is_active=eq.true&order=name&limit=1000' % SEDE)
chequear('Catalogo · getProducts devuelve', len(cat), len(maestro))
chequear('Catalogo · con su categoria resuelta', sum(1 for p in cat if p.get('categories')), len(maestro))
chequear('Catalogo · con los 5 niveles', sum(1 for p in cat if len(p.get('product_prices') or []) == 5), len(maestro))
chequear('Catalogo · con precio > 0', sum(1 for p in cat if p['price'] and float(p['price']) > 0), len(maestro))
# getCustomers(sedeId)
cli, _ = rest('customers', 'select=*&sede_id=eq.%s&is_active=eq.true&order=name&limit=1000' % SEDE)
chequear('Picker de clientes · getCustomers devuelve', len(cli), len(clientes))
chequear('Picker · sin nivel_default (a proposito)', sum(1 for c in cli if c['nivel_default'] is None), len(clientes))

print('\n══ RESULTADO ══')
if fallos:
    print('   🔴 %d verificaciones FALLARON:' % len(fallos))
    for f in fallos: print('      · ' + f)
    sys.exit(1)
print('   ✅ todas las verificaciones pasaron')
print('   ⛔ PENDIENTE y NO verificado aca: cartera (9 ventas a credito) e inventario inicial.')
